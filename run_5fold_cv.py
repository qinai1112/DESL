import argparse
import gc
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
import torch.optim as optim
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from sklearn.metrics import (
    average_precision_score,
    confusion_matrix,
    f1_score,
    mean_absolute_error,
    mean_squared_error,
    precision_score,
    recall_score,
    roc_auc_score,
)
from sklearn.model_selection import StratifiedKFold, StratifiedShuffleSplit
from torch_geometric.loader import DataLoader

import train as train_mod


METRIC_COLUMNS = [
    "Val AUC",
    "Val PR-AUC",
    "Val ACC",
    "Val Bal ACC",
    "Precision",
    "Recall",
    "F1",
    "Brier",
    "ECE",
]

FORMATTED_COLUMNS = [
    "Val AUC",
    "Val PR-AUC",
    "Val ACC",
    "Val Bal ACC",
    "Precision",
    "Recall",
    "Brier",
    "ECE",
]


def clone_graphs(graphs, indices):
    return [graphs[int(i)].clone() for i in indices]


def state_dict_with_ema(model, ema):
    """Materialize the EMA evaluation state used for model selection."""
    state = {}
    for key, value in model.state_dict().items():
        source = ema.shadow.get(key, value)
        state[key] = source.detach().cpu().clone()
    return state


def evaluate_silent(model, loader, device, threshold_metric="f1", threshold=None, return_probs=False):
    model.eval()
    y_true, y_prob = [], []
    with torch.no_grad():
        for batch in loader:
            batch = batch.to(device)
            out = model(batch)
            prob = torch.softmax(out, dim=1)[:, 1]
            y = batch.y.view(-1).long()
            y_true.extend(y.cpu().numpy().tolist())
            y_prob.extend(prob.cpu().numpy().tolist())

    try:
        auc = roc_auc_score(y_true, y_prob)
    except Exception:
        auc = float("nan")
    try:
        pr_auc = average_precision_score(y_true, y_prob)
    except Exception:
        pr_auc = float("nan")

    if threshold is None:  # validation: choose the threshold; test: apply the validation threshold
        threshold, _ = train_mod.find_best_threshold(y_true, y_prob, metric=threshold_metric)
    y_true_np = np.array(y_true)
    y_prob_np = np.array(y_prob)
    y_pred = (y_prob_np >= threshold).astype(int)

    precision = precision_score(y_true, y_pred, zero_division=0)
    recall = recall_score(y_true, y_pred, zero_division=0)
    f1 = f1_score(y_true, y_pred, zero_division=0)
    acc = (y_pred == y_true_np).mean()
    tp = ((y_true_np == 1) & (y_pred == 1)).sum()
    tn = ((y_true_np == 0) & (y_pred == 0)).sum()
    fp = ((y_true_np == 0) & (y_pred == 1)).sum()
    fn = ((y_true_np == 1) & (y_pred == 0)).sum()
    tpr = tp / (tp + fn + 1e-12)
    tnr = tn / (tn + fp + 1e-12)
    bal_acc = 0.5 * (tpr + tnr)
    mae = mean_absolute_error(y_true, y_prob)
    mse = mean_squared_error(y_true, y_prob)
    brier = mean_squared_error(y_true, y_prob)
    ece = train_mod.compute_ece(y_true, y_prob, n_bins=10)
    cm = confusion_matrix(y_true, y_pred)

    return {
        "Val AUC": float(auc),
        "Val PR-AUC": float(pr_auc),
        "Val ACC": float(acc),
        "Val Bal ACC": float(bal_acc),
        "Precision": float(precision),
        "Recall": float(recall),
        "F1": float(f1),
        "MAE": float(mae),
        "MSE": float(mse),
        "Brier": float(brier),
        "ECE": float(ece),
        "Thr": float(threshold),
        "CM": cm.tolist(),
        **({"_y_true": y_true, "_y_prob": y_prob} if return_probs else {}),
    }


def train_one_fold(
    fold,
    train_graphs,
    val_graphs,
    test_graphs,
    output_dir,
    epochs,
    lr,
    batch_size,
    patience,
    seed,
    best_by,
    threshold_metric,
    label_smoothing,
    use_class_weights,
    weight_decay,
    tab_dim,
    fusion_type,
    fusion_dim,
    image_dim=train_mod.IMAGE_FEATURE_DIM,
):
    train_mod.set_seed(seed)
    fold_dir = os.path.join(output_dir, f"fold_{fold}")
    os.makedirs(fold_dir, exist_ok=True)

    sample_in_dim = train_graphs[0].x.shape[1]
    tab_dim_for_data = sample_in_dim - image_dim
    if tab_dim_for_data != tab_dim:
        raise ValueError(
            f"Expected input_dim=image_dim+tab_dim ({image_dim}+{tab_dim}), got {sample_in_dim}"
        )
    if tab_dim_for_data % 2 != 0:
        raise ValueError(f"tab_dim must be even, got {tab_dim_for_data}")
    tab_dim_raw = tab_dim_for_data // 2
    train_mod.impute_missing_values(train_graphs, val_graphs, tab_dim_raw=tab_dim_raw)
    train_mod.impute_missing_values(train_graphs, test_graphs, tab_dim_raw=tab_dim_raw)  # 同样只用训练集中位数

    labels_arr = np.array([int(g.y.item()) for g in train_graphs], dtype=int)
    class_counts = np.bincount(labels_arr, minlength=2)
    if use_class_weights:
        class_weights = torch.tensor([1.0 / c if c > 0 else 0.0 for c in class_counts], dtype=torch.float)
        class_weights = class_weights / class_weights.sum() * len(class_counts)
    else:
        class_weights = None

    device = "cuda" if torch.cuda.is_available() else "cpu"
    if class_weights is not None:
        class_weights = class_weights.to(device)

    train_loader = DataLoader(train_graphs, batch_size=batch_size, shuffle=True)
    val_loader = DataLoader(val_graphs, batch_size=1, shuffle=False)
    test_loader = DataLoader(test_graphs, batch_size=1, shuffle=False)
    sample_batch = next(iter(train_loader))
    in_dim = sample_batch.x.size(-1)
    tab_dim_for_model = in_dim - image_dim

    model = train_mod.get_model(
        device,
        in_dim=in_dim,
        tab_dim=tab_dim_for_model,
        fusion_type=fusion_type,
        fusion_dim=fusion_dim,
    )
    loss_fn = train_mod.get_loss_fn(
        train_mod.default_loss_type,
        class_weights=class_weights,
        label_smoothing=label_smoothing,
    )
    optimizer = optim.AdamW(model.parameters(), lr=lr, weight_decay=weight_decay)
    scheduler = optim.lr_scheduler.ReduceLROnPlateau(optimizer, mode="max", factor=0.5, patience=2)
    ema = train_mod.EMA(model, decay=0.999)

    best_score, best_epoch, wait = -1.0, 0, 0
    best_row = None
    rows = []
    train_accs, val_accs = [], []
    model_path = os.path.join(fold_dir, "best_model.pt")

    for epoch in range(epochs):
        model.train()
        y_true_train, y_pred_train = [], []
        total_loss = 0.0

        for batch in train_loader:
            batch = batch.to(device)
            optimizer.zero_grad()
            out = model(batch)
            y = batch.y.view(-1).long()
            losses = loss_fn(out, y)
            losses["total"].backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), max_norm=5.0)
            optimizer.step()
            ema.update(model)

            total_loss += losses["total"].item()
            prob = torch.softmax(out, dim=1)[:, 1].detach().cpu().numpy()
            pred = (prob >= 0.5).astype(int)
            y_true_train.extend(y.detach().cpu().numpy().tolist())
            y_pred_train.extend(pred.tolist())

        train_acc = (np.array(y_true_train) == np.array(y_pred_train)).mean()
        train_accs.append(train_acc)

        backup = ema.swap_in(model)
        metrics = evaluate_silent(model, val_loader, device, threshold_metric=threshold_metric)
        ema.restore(model, backup)
        val_accs.append(metrics["Val ACC"])

        if best_by == "val_auc":
            val_score = metrics["Val AUC"]
        elif best_by == "val_bal_acc":
            val_score = metrics["Val Bal ACC"]
        else:
            val_score = metrics["Val PR-AUC"]
        scheduler.step(val_score)

        row = {
            "Fold": fold,
            "Epoch": epoch + 1,
            "Train ACC": float(train_acc),
            "Train Loss": float(total_loss / max(1, len(train_loader))),
            **{k: v for k, v in metrics.items() if k != "CM"},
            "CM": json.dumps(metrics["CM"]),
        }
        rows.append(row)
        pd.DataFrame(rows).to_csv(os.path.join(fold_dir, "training_log.csv"), index=False)

        print(
            f"[fold {fold}] epoch {epoch + 1}/{epochs} "
            f"train_acc={train_acc:.4f} auc={metrics['Val AUC']:.4f} "
            f"pr_auc={metrics['Val PR-AUC']:.4f} acc={metrics['Val ACC']:.4f} "
            f"f1={metrics['F1']:.4f} thr={metrics['Thr']:.2f}",
            flush=True,
        )

        if val_score > best_score:
            best_score, best_epoch, wait = val_score, epoch + 1, 0
            best_row = row.copy()
            torch.save(
                {
                    "model_state": state_dict_with_ema(model, ema),
                    "weights": "ema",
                    "threshold": float(metrics["Thr"]),
                    "best_by": best_by,
                    "best_score": float(val_score),
                    "fold": int(fold),
                    "tab_dim": int(tab_dim_for_model),
                    "fusion_type": fusion_type,
                    "fusion_dim": int(fusion_dim),
                },
                model_path,
            )
        else:
            wait += 1

        if wait >= patience:
            print(f"[fold {fold}] early stop at epoch {epoch + 1}; best epoch {best_epoch}", flush=True)
            break

    plt.figure(figsize=(6, 5))
    xs = range(1, len(val_accs) + 1)
    plt.plot(xs, train_accs, label="Train ACC")
    plt.plot(xs, val_accs, label="Val ACC")
    plt.xlabel("Epoch")
    plt.ylabel("Accuracy")
    plt.title(f"Fold {fold} Train vs Validation ACC")
    plt.legend()
    plt.grid(True)
    plt.tight_layout()
    plt.savefig(os.path.join(fold_dir, "train_val_curves.png"))
    plt.close()

    del model, optimizer, scheduler, ema, train_loader, val_loader
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()

    if best_row is None:
        raise RuntimeError(f"Fold {fold} did not produce a best row")

    # ---- Test on the held-out fold with the checkpoint and threshold selected on validation ----
    ckpt = torch.load(model_path, map_location=device, weights_only=False)
    test_model = train_mod.get_model(device, in_dim=in_dim, tab_dim=tab_dim_for_model,
                                     fusion_type=fusion_type, fusion_dim=fusion_dim)
    test_model.load_state_dict(ckpt["model_state"])
    test_metrics = evaluate_silent(test_model, test_loader, device, threshold=ckpt["threshold"], return_probs=True)
    pd.DataFrame({
        "patient_id": [str(getattr(g, "patient_id", "")) for g in test_graphs],
        "label": test_metrics.pop("_y_true"),
        "score": test_metrics.pop("_y_prob"),
    }).to_csv(os.path.join(fold_dir, "test_predictions.csv"), index=False)
    test_row = {
        "Fold": fold,
        "Best Epoch": best_epoch,
        "Inner Val AUC": best_row["Val AUC"],
        "Inner Val PR-AUC": best_row["Val PR-AUC"],
        **{k: v for k, v in test_metrics.items() if k != "CM"},  # held-out test metrics
        "CM": json.dumps(test_metrics["CM"]),
    }
    print(f"[fold {fold}] TEST auc={test_metrics['Val AUC']:.4f} pr_auc={test_metrics['Val PR-AUC']:.4f} "
          f"acc={test_metrics['Val ACC']:.4f} f1={test_metrics['F1']:.4f} thr(val)={test_metrics['Thr']:.2f}", flush=True)
    del test_model
    return test_row


def write_summary(best_rows, output_dir, method_name):
    fold_df = pd.DataFrame(best_rows)
    fold_path = os.path.join(output_dir, "fold_metrics.csv")
    fold_df.to_csv(fold_path, index=False)

    summary = {"Method": method_name}
    for col in METRIC_COLUMNS:
        summary[f"{col} Mean"] = float(fold_df[col].mean())
        summary[f"{col} Std"] = float(fold_df[col].std(ddof=0))
    summary_df = pd.DataFrame([summary])
    summary_df.to_csv(os.path.join(output_dir, "summary_metrics.csv"), index=False)

    formatted = {"Method": method_name}
    for col in FORMATTED_COLUMNS:
        mean = fold_df[col].mean()
        std = fold_df[col].std(ddof=0)
        formatted[col] = f"{mean:.4f}±{std:.4f}"
    formatted_df = pd.DataFrame([formatted])
    formatted_csv = os.path.join(output_dir, "summary_formatted.csv")
    formatted_xlsx = os.path.join(output_dir, "summary_formatted.xlsx")
    formatted_txt = os.path.join(output_dir, "summary_formatted.txt")
    formatted_df.to_csv(formatted_csv, index=False)
    formatted_df.to_excel(formatted_xlsx, index=False)
    with open(formatted_txt, "w", encoding="utf-8") as f:
        f.write("\t".join([formatted["Method"]] + [formatted[col] for col in FORMATTED_COLUMNS]) + "\n")

    wb = load_workbook(formatted_xlsx)
    ws = wb.active
    yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
    for cell in ws[2]:
        cell.fill = yellow
    wb.save(formatted_xlsx)

    print("\nFormatted result:", flush=True)
    print("\t".join([formatted["Method"]] + [formatted[col] for col in FORMATTED_COLUMNS]), flush=True)
    return fold_path, formatted_csv, formatted_xlsx, formatted_txt


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--feature_path",
        default=os.path.join(train_mod.DATASET_DIR, "features.npz"),
    )
    parser.add_argument(
        "--output_dir",
        default=os.path.join(train_mod.RESULTS_DIR, "next_visit_5fold"),
    )
    parser.add_argument("--method_name", default="Ours")
    parser.add_argument("--folds", type=int, default=5)
    parser.add_argument("--epochs", type=int, default=train_mod.DEFAULT_EPOCHS)
    parser.add_argument("--lr", type=float, default=train_mod.DEFAULT_LR)
    parser.add_argument("--batch_size", type=int, default=train_mod.DEFAULT_BATCH_SIZE)
    parser.add_argument("--patience", type=int, default=train_mod.DEFAULT_PATIENCE)
    parser.add_argument("--seed", type=int, default=train_mod.DEFAULT_SEED)
    parser.add_argument("--best_by", choices=["val_auc", "val_pr_auc", "val_bal_acc"], default="val_pr_auc")
    parser.add_argument("--threshold_metric", choices=["f1", "balanced_acc"], default="f1")
    parser.add_argument("--label_smoothing", type=float, default=0.0)
    parser.add_argument("--use_class_weights", type=lambda x: str(x).lower() == "true", default=True)
    parser.add_argument("--weight_decay", type=float, default=train_mod.DEFAULT_WEIGHT_DECAY)
    parser.add_argument("--tab_dim", type=int, default=train_mod.TAB_DIM)
    parser.add_argument(
        "--fusion_type",
        choices=["reliability_gate", "reliability_gate_no_interaction", "concat_mlp", "concat"],
        default="reliability_gate",
    )
    parser.add_argument("--fusion_dim", type=int, default=train_mod.DEFAULT_FUSION_DIM)
    args = parser.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)
    config = vars(args).copy()
    with open(os.path.join(args.output_dir, "config.json"), "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)

    graphs = train_mod.load_patient_graphs(args.feature_path)
    if not graphs:
        raise RuntimeError("No graphs built from feature_path")
    labels = np.array([int(g.y.item()) for g in graphs], dtype=int)
    print(
        f"Loaded graphs={len(graphs)} label_counts={dict(zip(*np.unique(labels, return_counts=True)))}",
        flush=True,
    )

    splitter = StratifiedKFold(n_splits=args.folds, shuffle=True, random_state=args.seed)
    best_rows = []
    for fold, (train_idx, val_idx) in enumerate(splitter.split(np.arange(len(graphs)), labels), start=1):
        # 2026-10-31: 留出折只作测试；其余四折再按患者 8:2 分出训练集和验证集
        inner = StratifiedShuffleSplit(n_splits=1, test_size=0.2, random_state=args.seed + fold)
        tr_rel, va_rel = next(inner.split(np.zeros(len(train_idx)), labels[train_idx]))
        train_graphs = clone_graphs(graphs, train_idx[tr_rel])
        val_graphs = clone_graphs(graphs, train_idx[va_rel])
        test_graphs = clone_graphs(graphs, val_idx)
        print(
            f"\n[fold {fold}] train={len(train_graphs)} val={len(val_graphs)} test={len(test_graphs)} "
            f"train_pos={sum(int(g.y.item()) for g in train_graphs)} val_pos={sum(int(g.y.item()) for g in val_graphs)} "
            f"test_pos={sum(int(g.y.item()) for g in test_graphs)}",
            flush=True,
        )
        best_row = train_one_fold(
            fold=fold,
            train_graphs=train_graphs,
            val_graphs=val_graphs,
            test_graphs=test_graphs,
            output_dir=args.output_dir,
            epochs=args.epochs,
            lr=args.lr,
            batch_size=args.batch_size,
            patience=args.patience,
            seed=args.seed + fold,
            best_by=args.best_by,
            threshold_metric=args.threshold_metric,
            label_smoothing=args.label_smoothing,
            use_class_weights=args.use_class_weights,
            weight_decay=args.weight_decay,
            tab_dim=args.tab_dim,
            fusion_type=args.fusion_type,
            fusion_dim=args.fusion_dim,
        )
        best_rows.append(best_row)
        pd.DataFrame(best_rows).to_csv(os.path.join(args.output_dir, "fold_metrics_partial.csv"), index=False)

    write_summary(best_rows, args.output_dir, args.method_name)


if __name__ == "__main__":
    main()
